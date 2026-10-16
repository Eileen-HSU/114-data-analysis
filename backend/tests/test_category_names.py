#!/usr/bin/env python
"""
大類別 / 子類別名稱的顯示語言（services/category_name_service.py）。

涵蓋：
    1. 中文介面：任何輸入都原樣（輸出與改動前相同）、不呼叫 AI
    2. 英文介面：內建 4 個大類別 + 21 個子類別全部有英文、不含漢字、代碼前綴保留、不呼叫 AI
    3. 自訂 / AI 新類別：批次一次呼叫、有快取（第二次不再呼叫）、保留代碼前綴
    4. 失敗保護：AI 例外、回傳含漢字、空字串、格式錯誤 -> 退回原名稱，不拋例外，下次可重試
    5. 匯出（路由）：英文介面的 Excel / Word 整份沒有漢字；中文介面維持原名稱；分組合併仍正確
    6. 翻譯 API：需登入、參數驗證、依 Accept-Language

執行方式：
    cd backend
    python3 tests/test_category_names.py
"""

import base64
import io
import re

import docx
import openpyxl

from admin_test_support import GEMINI_CALLS, GEMINI_QUEUE, check, create_app, finish, q, seed_people, user_header
import models as m
from extensions import db
from services import category_name_service as C
from services.subcategory_methodology import SUBCATEGORY_METHODOLOGY

HAN = re.compile(r"[一-鿿]")


def reset():
    GEMINI_QUEUE.clear(); GEMINI_CALLS.clear(); C._reset_cache_for_tests()


print("========== 1. 中文介面原樣 ==========")
reset()
names = ["主管領導", "A1 工作與生活邊界", "自訂類別", "X"]
check("zh-TW 原樣", C.translate_category_names(names, "zh-TW") == {n: n for n in names})
check("lang=None 原樣", C.translate_category_names(names, None) == {n: n for n in names})
check("lang 不合法原樣", C.translate_category_names(names, "fr") == {n: n for n in names})
rows = [{"main_category": "主管領導", "sub_category": "A1 工作與生活邊界"}]
check("匯出 rows：zh-TW 回傳同一份（不複製、不變）", C.localize_category_rows(rows, "zh-TW") is rows)
check("中文介面不呼叫 AI", not GEMINI_CALLS)

print("\n========== 2. 內建分類全部有英文 ==========")
reset()
mains = sorted({v["main_category"] for t in SUBCATEGORY_METHODOLOGY.values() for v in t.values()})
subs = sorted({k for t in SUBCATEGORY_METHODOLOGY.values() for k in t})
res = C.translate_category_names(mains + subs, "en")
check(f"內建大類別 {len(mains)} 個 + 子類別 {len(subs)} 個：全部翻成英文、無漢字", all(not HAN.search(v) for v in res.values()))
check("代碼前綴保留（A1 Work-Life Boundaries）", res["A1 工作與生活邊界"] == "A1 Work-Life Boundaries")
check("大類別翻譯", res["主管領導"] == "Supervisor Leadership")
check("內建分類完全不呼叫 AI", not GEMINI_CALLS)
check("已是英文的名稱原樣", C.translate_category_names(["Learning", "B2 Equipment"], "en") == {"Learning": "Learning", "B2 Equipment": "B2 Equipment"})

print("\n========== 3. 自訂 / AI 新類別 ==========")
reset()
q({"薪資福利": "Compensation & Benefits", "遠距工作": "Remote Work"})
res = C.translate_category_names(["薪資福利", "C3 遠距工作", "主管領導", "薪資福利"], "en")
check("自訂類別翻成英文", res["薪資福利"] == "Compensation & Benefits")
check("自訂類別保留代碼前綴", res["C3 遠距工作"] == "C3 Remote Work")
check("內建類別不送 AI（只送兩個自訂名稱、重複只送一次）", len(GEMINI_CALLS) == 1 and "主管領導" not in GEMINI_CALLS[0]["contents"] and GEMINI_CALLS[0]["contents"].count("薪資福利") == 1)
GEMINI_CALLS.clear()
res2 = C.translate_category_names(["薪資福利", "C3 遠距工作"], "en")
check("第二次走快取、不再呼叫 AI", res2["薪資福利"] == "Compensation & Benefits" and not GEMINI_CALLS)

print("\n========== 4. 失敗保護 ==========")
for label, reply in (("AI 例外", RuntimeError("boom")), ("仍含漢字", {"職場文化": "職場 culture"}), ("空字串", {"職場文化": "  "}),
                     ("格式錯誤（陣列）", ["a"]), ("不是 JSON", "not json"), ("值不是字串", {"職場文化": 5})):
    reset(); q(reply)
    try:
        r = C.translate_category_names(["職場文化"], "en")
        check(f"{label}：退回原名稱、不拋例外", r == {"職場文化": "職場文化"})
    except Exception as e:  # noqa: BLE001
        check(f"{label}：退回原名稱、不拋例外（拋了 {e!r}）", False)
reset(); q(RuntimeError("boom")); C.translate_category_names(["職場文化"], "en")
q({"職場文化": "Workplace Culture"})
check("失敗不快取：下次可重試成功", C.translate_category_names(["職場文化"], "en")["職場文化"] == "Workplace Culture")
reset(); q({"別的": "Other"})
check("AI 回傳不相干的 key 被忽略", C.translate_category_names(["職場文化"], "en") == {"職場文化": "職場文化"})
check("None / 非字串名稱不炸", C.translate_category_names([None, 5, "", "主管領導"], "en") == {"主管領導": "Supervisor Leadership"})

print("\n========== 5. 匯出路由 ==========")
reset()
app = create_app(); client = app.test_client()
with app.app_context():
    seed_people()
    if db.session.get(m.Workspace, 1) is None:
        db.session.add(m.Workspace(project_id=1, user_id=1, project_name="ws")); db.session.flush()
    chat = m.Chat_History(project_id=1, sender_type="ai", message_content="legacy text, not a classification snapshot")
    db.session.add(chat); db.session.commit(); chat_id = chat.chat_id

ROWS = [
    {"main_category": "主管領導", "sub_category": "A1 工作與生活邊界", "respondent_text": "受試者1：I get messages after hours",
     "aggregated_reasoning": "Boundary issue.", "aggregated_summary": "Respect off hours.", "source_column": "Q1"},
    {"main_category": "主管領導", "sub_category": "A2 回饋與溝通", "respondent_text": "受試者2：More feedback please",
     "aggregated_reasoning": "Needs feedback.", "aggregated_summary": "Give feedback.", "source_column": "Q1"},
    {"main_category": "部門合作", "sub_category": "B1 溝通與協調機制", "respondent_text": "受試者3：Meetings overlap",
     "aggregated_reasoning": "Coordination.", "aggregated_summary": "Align calendars.", "source_column": "Q1"},
]


def do_export(lang, kind):
    h = dict(user_header(1))
    if lang:
        h["Accept-Language"] = lang
    r = client.post("/api/exports", json={"export_type": kind, "chat_id": chat_id, "filename": "r", "rows": ROWS}, headers=h)
    assert r.status_code in (200, 201), r.get_data(as_text=True)
    with app.app_context():
        return base64.b64decode(db.session.get(m.Export_File, r.get_json()["export_id"]).content)


def xlsx_all(data):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value not in (None, "")] + [ws.title for ws in wb.worksheets]


def docx_all(data):
    d = docx.Document(io.BytesIO(data))
    return [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]


en_x = xlsx_all(do_export("en", "xlsx")); en_d = docx_all(do_export("en", "docx"))
check("英文 Excel：整份沒有漢字（含類別名稱）", not [t for t in en_x if HAN.search(t)])
check("英文 Word：整份沒有漢字（含類別名稱）", not [t for t in en_d if HAN.search(t)])
check("英文 Excel：看得到英文類別名稱", "Supervisor Leadership" in en_x and "A1 Work-Life Boundaries" in en_x and "Departmental Collaboration" in en_x)
zh_x = xlsx_all(do_export("zh-TW", "xlsx")); none_x = xlsx_all(do_export(None, "xlsx"))
check("中文 Excel：類別名稱維持原樣", "主管領導" in zh_x and "A1 工作與生活邊界" in zh_x and "Supervisor Leadership" not in zh_x)
check("沒有 Accept-Language -> 與中文相同", none_x == zh_x)
wb = openpyxl.load_workbook(io.BytesIO(do_export("en", "xlsx")))
ws = wb.worksheets[0]
merged = [str(r) for r in ws.merged_cells.ranges]
check("英文 Excel：同一大類別仍合併儲存格（A2:A3），值是英文", "A2:A3" in merged and ws["A2"].value == "Supervisor Leadership" and ws["A4"].value == "Departmental Collaboration")
check("匯出不改動資料庫裡的原始名稱（ROWS 原樣）", ROWS[0]["main_category"] == "主管領導")

print("\n========== 6. 翻譯 API ==========")
reset()
h_en = {**user_header(1), "Accept-Language": "en"}
check("未登入 -> 401", client.post("/api/classification/category-names", json={"names": ["主管領導"]}, headers={"Accept-Language": "en"}).status_code == 401)
check("names 不是陣列 -> 400", client.post("/api/classification/category-names", json={"names": "x"}, headers=h_en).status_code == 400)
check("names 超過 500 筆 -> 400", client.post("/api/classification/category-names", json={"names": ["a"] * 501}, headers=h_en).status_code == 400)
r = client.post("/api/classification/category-names", json={"names": ["主管領導", "A1 工作與生活邊界"]}, headers=h_en)
check("英文：回傳英文名稱", r.status_code == 200 and r.get_json()["names"] == {"主管領導": "Supervisor Leadership", "A1 工作與生活邊界": "A1 Work-Life Boundaries"})
r = client.post("/api/classification/category-names", json={"names": ["主管領導"]}, headers={**user_header(1), "Accept-Language": "zh-TW"})
check("中文：原樣回傳", r.get_json()["names"] == {"主管領導": "主管領導"})
r = client.post("/api/classification/category-names", json={"names": ["主管領導"]}, headers=user_header(1))
check("沒有 Accept-Language：原樣回傳", r.get_json()["names"] == {"主管領導": "主管領導"})

print("\n========== 7. 翻譯專用 key ==========")
import os
import services.gemini_client as GC
seen_keys = []
_Fake = GC.GenerativeModel


class KeyRecorder(_Fake):
    def generate_content(self, contents, **kw):
        seen_keys.append(GC.current_api_key())
        return super().generate_content(contents, **kw)


GC.GenerativeModel = KeyRecorder
GC._api_key = "main-key"
os.environ.pop("TRANSLATION_GEMINI_API_KEY", None)
reset(); q({"職場文化": "Workplace Culture"}); seen_keys.clear()
C.translate_category_names(["職場文化"], "en")
check("沒設定專用 key：沿用原本的 key", seen_keys == ["main-key"])
os.environ["TRANSLATION_GEMINI_API_KEY"] = "translate-key"
reset(); q({"職場文化": "Workplace Culture"}); seen_keys.clear()
C.translate_category_names(["職場文化"], "en")
check("設定專用 key：翻譯用專用 key", seen_keys == ["translate-key"])
check("翻譯結束後 key 還原（不影響後續的分類 / 摘要）", GC.current_api_key() == "main-key")
reset(); q(RuntimeError("quota")); seen_keys.clear()
C.translate_category_names(["職場文化"], "en")
check("翻譯失敗時 key 也會還原", GC.current_api_key() == "main-key")
os.environ.pop("TRANSLATION_GEMINI_API_KEY", None)
GC.GenerativeModel = _Fake

print("\n========== 8. AI 文字用的類別名稱 ==========")
from services import aggregated_summary_service as AS
OK = {"reasoning_summary": "r", "summary_summary": "s"}
reset(); q(OK)
AS.build_aggregated_summary_pair("主管領導", "A1 工作與生活邊界", [{"matched_segment_text": "x"}], [{"matched_segment_text": "y"}], lang="en")
c = GEMINI_CALLS[-1]["contents"]
check("英文摘要：給 AI 的類別是英文名稱", "Supervisor Leadership" in c and "A1 Work-Life Boundaries" in c and "主管領導" not in c)
reset(); q(OK)
AS.build_aggregated_summary_pair("主管領導", "A1 工作與生活邊界", [{"matched_segment_text": "x"}], [{"matched_segment_text": "y"}], lang="zh-TW")
c = GEMINI_CALLS[-1]["contents"]
check("中文摘要：類別維持中文（與改動前相同）", "類別：主管領導 / A1 工作與生活邊界" in c)
reset(); q({"summary": "ok"})
AS.build_aggregated_summary("主管領導", "A1 工作與生活邊界", [{"matched_segment_text": "x"}], lang="en")
check("單組摘要（英文）也用英文名稱", "類別：Supervisor Leadership / A1 Work-Life Boundaries" in GEMINI_CALLS[-1]["contents"])

from services import chat_ask_service as CA
ITEMS = [{"main_category": "主管領導", "sub_category": "A1 工作與生活邊界", "respondent_number": 1, "status": "completed",
          "excerpt": "e", "reasoning": "r", "summary": "s", "methodology": None, "citation": None,
          "secondary_categories": [{"main_category": "部門合作", "sub_category": "B1 溝通與協調機制"}]}]
agg = {("主管領導", "A1 工作與生活邊界"): {"aggregated_reasoning": "ar", "aggregated_summary": "as"}}
en_ctx = CA._build_context_text(ITEMS, "Summarize everything", agg, "en")
check("Chat 追問（英文）：context 的類別是英文，含次要分類", "Supervisor Leadership／A1 Work-Life Boundaries" in en_ctx and "Departmental Collaboration／B1 Communication & Coordination Mechanisms" in en_ctx and "主管領導" not in en_ctx)
zh_ctx = CA._build_context_text(ITEMS, "整體有哪些發現", agg, "zh-TW")
none_ctx = CA._build_context_text(ITEMS, "整體有哪些發現", agg)
check("Chat 追問（中文 / 未指定）：維持中文，兩者相同", "主管領導／A1 工作與生活邊界" in zh_ctx and zh_ctx == none_ctx)
code_ctx = CA._build_context_text(ITEMS, "Tell me about A1", agg, "en")
check("英文介面用代碼（A1）查詢仍比對得到（比對用原名稱）", "[受試者1]" in code_ctx and "沒有找到符合的片段" not in code_ctx)

finish()
