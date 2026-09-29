#!/usr/bin/env python
"""
問卷新增填答後，Chat 分類結果訊息的 rating_stats 不可以停在分析當下的快照：
freshness 判定過期 -> refresh 重算 -> Excel / Word 匯出跟畫面數字一致。

涵蓋：
    1. 分析當下存進 Chat_History 的 rating_stats（1 份填答）
    2. 新增一份只有 rating 的填答：freshness stale=True（rating_stats_stale）
    3. refresh：rating_stats 重算（2 份、平均更新）、寫回 Chat_History、
       分類 rows 不受影響；refresh 後 freshness 不再過期
    4. 匯出 xlsx / docx：即使前端送來舊的 rating_stats，也以後端最新版本為準，
       跟 refresh 後畫面數字一致
    5. rating 題仍然獨立統計（不進 Response_Classification）

執行方式：
    cd backend
    python3 tests/test_rating_stats_freshness.py
"""

import base64
import io
import os

import openpyxl
from docx import Document

from admin_test_support import GEMINI_QUEUE, check, create_app, finish, q, seed_people, seed_topic, user_header
import models as m
from extensions import db
from services import workspace_result_service as wrs
from services.privacy_service import mask_pii

os.environ.pop("OPEN_CLASSIFICATION_ENABLED", None)
app = create_app()
client = app.test_client()

with app.app_context():
    seed_people()
    seed_topic("career", categories=[("職涯發展", "A1 教育訓練", "m", "c")])
    tpl = m.Survey_Template(user_id=1, title="滿意度", access_code="RATE1", question_json={"items": [
        {"id": "r1", "type": "rating", "title": "整體滿意度"},
        {"id": "q1", "type": "short", "title": "建議", "question_type": "career", "routing_status": "routed"},
    ]})
    db.session.add(tpl)
    db.session.flush()
    TEMPLATE_ID = tpl.template_id
    db.session.add(m.Survey_Response(template_id=TEMPLATE_ID, answer_json={"answers": {"r1": 4, "q1": "希望多開課"}}))
    db.session.add(m.Workspace(project_id=7, user_id=1, project_name="ws"))
    db.session.commit()

GEMINI_QUEUE.clear()
q({"segments": [mask_pii("希望多開課")]})
q({"classifications": [{"index": 0, "main_category": "職涯發展", "sub_category": "A1 教育訓練",
                         "secondary_sub_category": None, "reasoning": "r", "summary": "s", "confidence": 0.9}]})
analyze = client.post("/api/surveys/RATE1/analyze", headers=user_header(1)).get_json()
GEMINI_QUEUE.clear()
snapshot_stats = analyze["rating_stats"]
check("分析當下：1 份 rating、平均 4.0", snapshot_stats[0]["answered_count"] == 1 and snapshot_stats[0]["average"] == 4.0)

with app.app_context():
    chat = m.Chat_History(project_id=7, template_id=TEMPLATE_ID, sender_type="ai", message_content=wrs.build_classification_message({
        "rows": [{"main_category": "職涯發展", "sub_category": "A1 教育訓練", "count": 1}],
        "meta": {"template_id": TEMPLATE_ID, "source_type": "survey", "review_revision": analyze["review_revision"]},
        "rating_stats": snapshot_stats,
    }))
    db.session.add(chat)
    db.session.commit()
    CHAT_ID = chat.chat_id

fresh = client.get(f"/api/chat/{CHAT_ID}/classification-result/freshness", headers=user_header(1)).get_json()
check("還沒有新填答：不過期", fresh["stale"] is False)

with app.app_context():
    db.session.add(m.Survey_Response(template_id=TEMPLATE_ID, answer_json={"answers": {"r1": 1}}))
    db.session.commit()

fresh = client.get(f"/api/chat/{CHAT_ID}/classification-result/freshness", headers=user_header(1)).get_json()
check("新增填答後：freshness stale（rating_stats_stale）", fresh["stale"] is True and fresh.get("rating_stats_stale") is True)

resp = client.post(f"/api/chat/{CHAT_ID}/classification-result/refresh", headers=user_header(1))
body = resp.get_json()
check("refresh 200", resp.status_code == 200)
live = body["rating_stats"][0]
check("refresh 後 rating_stats 重算：2 份、平均 2.5、分布正確",
      live["answered_count"] == 2 and live["average"] == 2.5 and live["distribution"]["1"] == 1 and live["distribution"]["4"] == 1)
check("分類 rows 沒有被 rating 影響", len(body["rows"]) == 1)
with app.app_context():
    stored = wrs.parse_classification_message(db.session.get(m.Chat_History, CHAT_ID).message_content)
    check("Chat_History 已寫回最新 rating_stats", stored["rating_stats"][0]["answered_count"] == 2)
    check("rating 題沒有進 Response_Classification", m.Response_Classification.query.filter_by(question_id="r1").count() == 0)
fresh = client.get(f"/api/chat/{CHAT_ID}/classification-result/freshness", headers=user_header(1)).get_json()
check("refresh 後不再過期", fresh["stale"] is False)

# 再新增一份填答（畫面尚未 refresh），前端帶著舊快照匯出
with app.app_context():
    db.session.add(m.Survey_Response(template_id=TEMPLATE_ID, answer_json={"answers": {"r1": 1}}))
    db.session.commit()


def export(kind):
    resp = client.post("/api/exports", headers=user_header(1), json={
        "chat_id": CHAT_ID, "filename": f"out.{kind}", "export_type": kind, "rows": [],
        "rating_stats": snapshot_stats,  # 前端可能還是舊快照
    })
    data = resp.get_json()
    with app.app_context():
        content = db.session.get(m.Export_File, data["export_id"]).content
    return resp.status_code, base64.b64decode(content)


status, xlsx_bytes = export("xlsx")
wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
sheet_text = " ".join(str(c.value) for row in wb["評分題統計"].iter_rows() for c in row if c.value is not None)
check("xlsx 匯出 201", status == 201)
check("xlsx 用最新的 3 份填答（不是前端舊快照的 1 份）", "3" in sheet_text and "2.0" in sheet_text)
status, docx_bytes = export("docx")
doc_text = "\n".join(p.text for p in Document(io.BytesIO(docx_bytes)).paragraphs) + "\n" + "\n".join(
    cell.text for t in Document(io.BytesIO(docx_bytes)).tables for r in t.rows for cell in r.cells)
check("docx 匯出 201", status == 201)
check("docx 用最新平均 2.0", "2.0" in doc_text)
with app.app_context():
    stored = wrs.parse_classification_message(db.session.get(m.Chat_History, CHAT_ID).message_content)
    check("匯出時也把 Chat_History 更新成同一份數字（畫面重新整理後一致）",
          stored["rating_stats"][0]["answered_count"] == 3 and stored["rating_stats"][0]["average"] == 2.0)

finish()
