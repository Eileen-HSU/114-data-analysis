import unittest
from io import BytesIO
from types import SimpleNamespace

from docx import Document
import pandas as pd
from services.classification_upload import read_analysis_upload, UploadValidationError, MAX_UPLOAD_BYTES


def upload(name, content):
    return SimpleNamespace(filename=name, stream=BytesIO(content))


class UploadReaderTests(unittest.TestCase):
    def test_excel(self):
        buffer = BytesIO()
        pd.DataFrame({"feedback": ["Useful course"]}).to_excel(buffer, index=False)
        self.assertEqual(read_analysis_upload(upload("data.XLSX", buffer.getvalue())).iloc[0, 0], "Useful course")

    def test_word_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("First answer")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Table answer"
        document.add_paragraph("Last answer")
        buffer = BytesIO()
        document.save(buffer)
        self.assertEqual(read_analysis_upload(upload("data.docx", buffer.getvalue()))["Word"].tolist(), ["First answer", "Table answer", "Last answer"])

    def test_size_limit_before_parsing(self):
        with self.assertRaises(UploadValidationError) as error:
            read_analysis_upload(upload("data.docx", b"x" * (MAX_UPLOAD_BYTES + 1)))
        self.assertEqual(error.exception.status, 413)

    def test_exact_limit_is_not_rejected_for_size(self):
        with self.assertRaises(UploadValidationError) as error:
            read_analysis_upload(upload("data.xlsx", b"x" * MAX_UPLOAD_BYTES))
        self.assertEqual(error.exception.status, 400)

    def test_unsupported_or_corrupt_files(self):
        for name in ("data.csv", "data.txt", "data.doc", "data.docx", "data.xlsx"):
            with self.subTest(name=name), self.assertRaises(UploadValidationError):
                read_analysis_upload(upload(name, b"invalid"))

    def test_empty_word(self):
        buffer = BytesIO()
        Document().save(buffer)
        with self.assertRaises(UploadValidationError):
            read_analysis_upload(upload("empty.docx", buffer.getvalue()))


if __name__ == "__main__":
    unittest.main()
