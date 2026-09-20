#!/usr/bin/env python
"""
測試腳本：問卷「Chat 分析」新增的評分題統計功能。

涵蓋範圍：
    Part A - _safe_rating_int() 單元測試（合法/非法 rating 值判斷）
    Part B - _build_rating_stats() 單元測試（純函式，不碰資料庫）
        - rating=0 正確計入平均/answered_count/distribution
        - rating=5 正確計入
        - 未作答（key 不存在）不計入任何統計
        - 非法值（超出範圍／無法轉換）直接跳過
        - 完全沒人作答的 rating 題：average=None、answered_count=0，
          distribution 六桶皆 0，但仍要出現在結果陣列裡
        - question_number 對應 items 原始位置（含跟 short 題混合的情況）
        - title / question_title 兩種欄位都要能讀到
        - 問卷沒有 rating 題 -> rating_stats == []
    Part C - POST /api/surveys/<access_code>/analyze 端到端測試
        - rating-only 問卷（無 short 題）：早退分支也要帶 rating_stats
        - short-only 問卷（有 Published Taxonomy，真的跑分類）：
          rating_stats == []，既有分類行為不受影響（regression）
        - rating + short 混合、taxonomy 可用：兩邊互不干擾，各自正確
        - rating + short 混合、taxonomy 不可用：rating_stats 仍正確算出，
          不因為分類那邊 taxonomy_unavailable 而被拖累
    Part D - build_xlsx() / build_docx() 的 rating_stats 可選參數
        - 有 rating_stats：Excel 多一張「評分題統計」sheet（rows=[] 也不能壞）、
          Word 在既有分類結果之前多一段「評分題統計」
        - 沒有 rating_stats（None / 未傳）：輸出跟這個參數新增之前
          完全一樣（regression）
    Part E - POST /api/exports：rating-only（rows=[] 但有 rating_stats）
          必須成功，不能因為 rows 是空陣列就被擋掉

執行方式：
    cd backend
    python3 test_survey_rating_stats.py
"""

import io
import json
import os
import sys
from datetime import datetime, timezone

os.environ.setdefault("JWT_SECRET_KEY", "test-secret-key-for-testing-only")
sys.path.insert(0, os.path.dirname(__file__))

FAILED = []


def check(label, condition):
    status = "PASS" if condition else "FAIL"
    if status == "FAIL":
        FAILED.append(label)
    print(f"[{status}] {label}")


# ── 假的 Gemini：monkeypatch services.gemini_client.GenerativeModel ──
# 不去動 sys.modules["google"]／["google.genai"]：這個專案的正式呼叫路徑
# 已經統一收斂到 services/gemini_client.py 這一層薄封裝（内部才是
# `from google import genai`），classify_v2.py／aggregated_summary_service.py
# 都是透過 `from services import gemini_client as genai` 再呼叫
# `genai.GenerativeModel(...)`，所以只要在這一層換掉 GenerativeModel
# 類別本身，就能讓分類流程吃到假回應，同時完全不影響真正的
# `google-genai` SDK（環境裡如果沒裝也沒關係，這裡從頭到尾不會真的建立
# genai.Client()）。
import services.gemini_client as gemini_client

_queue = []


def q(obj_or_text):
    _queue.append(obj_or_text if isinstance(obj_or_text, str) else json.dumps(obj_or_text, ensure_ascii=False))


class _FakeResp:
    def __init__(self, text):
        self.text = text


class _FakeGenerativeModel:
    def __init__(self, model_name=None, system_instruction=None, **kwargs):
        pass

    def generate_content(self, contents, **kwargs):
        return _FakeResp(_queue.pop(0))


gemini_client.GenerativeModel = _FakeGenerativeModel


import jwt
from flask import Flask
from sqlalchemy.dialects.mysql import MEDIUMTEXT
from sqlalchemy.ext.compiler import compiles

from extensions import db
import models as m

# ── SQLite 相容性 shim：MEDIUMTEXT 是 MySQL 專屬型別，這個環境裝的
# SQLAlchemy 版本在 SQLite 上編譯 CREATE TABLE 時不會自動 fallback 成
# TEXT（其他既有測試檔案，例如 test_workspace_sharing.py，在同一個環境
# 下建 Chat_History 表也會撞到一模一樣的 CompileError，屬於測試環境既有
# 限制，不是這次改動造成的）。這裡只在 sqlite 方言下把 MEDIUMTEXT 編譯
# 成 TEXT，只影響「這支測試檔案自己開的 SQLite 記憶體資料庫」，完全不
# 碰 models.py、不影響正式環境（正式環境是 MySQL，本來就有 MEDIUMTEXT）。
@compiles(MEDIUMTEXT, "sqlite")
def _compile_mediumtext_sqlite(element, compiler, **kw):
    return "TEXT"
from services.privacy_service import mask_pii
from services.subcategory_methodology import QUESTION_LEADERSHIP, SUBCATEGORY_METHODOLOGY
from routes.classifications.classification import (
    classification_bp,
    _safe_rating_int,
    _build_rating_stats,
)
from routes.surveys.survey import survey_bp
from routes.exports.export import exports_bp
from services.export_file_service import (
    build_xlsx,
    build_docx,
    COLUMN_HEADERS,
    _render_rating_donut_png,
)
import openpyxl
from docx import Document as _DocxDocument

app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///:memory:"
app.config["TESTING"] = True
app.register_blueprint(classification_bp)
app.register_blueprint(survey_bp)
app.register_blueprint(exports_bp)
db.init_app(app)

with app.app_context():
    db.metadata.create_all(
        bind=db.engine,
        tables=[
            m.User.__table__,
            m.Survey_Template.__table__,
            m.Survey_Response.__table__,
            m.Response_Classification.__table__,
            m.Response_Segmentation_Status.__table__,
            m.Uploaded_Answer.__table__,
            m.Topic.__table__,
            m.Taxonomy_Version.__table__,
            m.Taxonomy_Category.__table__,
            m.Workspace.__table__,
            m.Chat_History.__table__,
            m.Export_File.__table__,
        ],
    )

    # 建一份 leadership_and_dept 的 Published Taxonomy，讓「short 題」在
    # 測試裡可以真的走完整套分類（而不是每次都卡在 taxonomy_unavailable），
    # 比照 test_batch_classification.py 既有的作法。
    db.session.add(m.Topic(topic_key=QUESTION_LEADERSHIP, title="主管領導和部門合作"))
    version = m.Taxonomy_Version(
        topic_key=QUESTION_LEADERSHIP, version_number=1,
        status="published", source="migrated_legacy",
    )
    db.session.add(version)
    db.session.flush()
    for i, (sub_category, info) in enumerate(SUBCATEGORY_METHODOLOGY[QUESTION_LEADERSHIP].items(), start=1):
        db.session.add(m.Taxonomy_Category(
            version_id=version.version_id,
            main_category=info["main_category"],
            sub_category=sub_category,
            methodology=info["methodology"],
            citation=info["citation"],
            source_raw_text=f"{sub_category}：（測試用簡化規則原文）",
            sort_order=i,
        ))

    db.session.add(m.User(user_id=1, user_name="owner", email="owner@example.com", password_hash="x"))
    db.session.add(m.User(user_id=2, user_name="stranger", email="stranger@example.com", password_hash="x"))
    db.session.commit()

client = app.test_client()


def auth_header(user_id):
    token = jwt.encode({"user_id": user_id}, os.environ["JWT_SECRET_KEY"], algorithm="HS256")
    return {"Authorization": f"Bearer {token}"}


def make_survey(*, access_code, items, owner_id=1):
    with app.app_context():
        survey = m.Survey_Template(
            title=f"測試問卷_{access_code}",
            access_code=access_code,
            user_id=owner_id,
            question_json={"description": "", "identity_mode": "anonymous", "items": items},
        )
        db.session.add(survey)
        db.session.commit()
        return survey.template_id


def add_response(template_id, answers):
    with app.app_context():
        response = m.Survey_Response(template_id=template_id, answer_json={"answers": answers})
        db.session.add(response)
        db.session.commit()
        return response.response_id


# 提供給 Part B 用的輕量假 response 物件，不需要真的建進資料庫
class FakeResponse:
    def __init__(self, answers):
        self.answer_json = {"answers": answers}


def read_xlsx_sheets(file_bytes):
    """回傳 {sheet_title: [[cell,...], ...]}，方便斷言每張 sheet 的內容。"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    return {ws.title: [[cell.value for cell in row] for row in ws.iter_rows()] for ws in wb.worksheets}, wb.sheetnames


def read_docx_text(file_bytes):
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


# ═══════════════════════════════════════════════════════════════
# Part A：_safe_rating_int() 單元測試
# ═══════════════════════════════════════════════════════════════
print("========== Part A：_safe_rating_int() ==========")

check("字串 '0' -> 0", _safe_rating_int("0") == 0)
check("字串 '5' -> 5", _safe_rating_int("5") == 5)
check("整數 0 -> 0", _safe_rating_int(0) == 0)
check("整數 5 -> 5", _safe_rating_int(5) == 5)
check("浮點數 3.0 -> 3", _safe_rating_int(3.0) == 3)
check("超出上限 '6' -> None", _safe_rating_int("6") is None)
check("負數 '-1' -> None", _safe_rating_int("-1") is None)
check("非數字字串 'abc' -> None", _safe_rating_int("abc") is None)
check("None -> None", _safe_rating_int(None) is None)
check("空字串 '' -> None", _safe_rating_int("") is None)
check("bool True 不會被誤轉成 1", _safe_rating_int(True) is None)
check("bool False 不會被誤轉成 0", _safe_rating_int(False) is None)
check("list -> None（型別完全不對）", _safe_rating_int([0]) is None)


# ═══════════════════════════════════════════════════════════════
# Part B：_build_rating_stats() 單元測試（純函式，不碰資料庫）
# ═══════════════════════════════════════════════════════════════
print("\n========== Part B：_build_rating_stats() ==========")

items_b1 = [
    {"id": "q_short_0", "type": "short", "title": "排在最前面的開放題"},
    {"id": "q_rating", "type": "rating", "title": "課程整體滿意度"},
    {"id": "q_short_2", "type": "short", "title": "排在後面的開放題"},
]
responses_b1 = [
    FakeResponse({"q_rating": "0"}),               # 合法：0 分
    FakeResponse({"q_rating": "5"}),                # 合法：5 分
    FakeResponse({"q_rating": "3"}),                # 合法：3 分
    FakeResponse({"q_short_0": "沒有回答評分題"}),   # 未作答：key 不存在
    FakeResponse({"q_rating": "9"}),                # 非法：超出範圍，跳過
    FakeResponse({"q_rating": "not-a-number"}),     # 非法：無法轉換，跳過
]
stats_b1 = _build_rating_stats(items_b1, responses_b1)

check("只有 1 題 rating，結果陣列長度為 1", len(stats_b1) == 1)
stat_b1 = stats_b1[0] if stats_b1 else {}
check("question_id 正確", stat_b1.get("question_id") == "q_rating")
check(
    "question_number = 該題在 items 的原始位置 + 1（排第 2 個，含前面的 short 題）",
    stat_b1.get("question_number") == 2,
)
check("title 正確", stat_b1.get("title") == "課程整體滿意度")
check("rating=0 有被計入 answered_count（3 筆合法：0,5,3）", stat_b1.get("answered_count") == 3)
check("平均分只用 3 筆合法值計算：(0+5+3)/3 = 2.7", stat_b1.get("average") == 2.7)
check("distribution 裡 0 分人數是 1（不是被排除）", stat_b1.get("distribution", {}).get("0") == 1)
check("distribution 裡 5 分人數是 1", stat_b1.get("distribution", {}).get("5") == 1)
check("distribution 裡 3 分人數是 1", stat_b1.get("distribution", {}).get("3") == 1)
check("distribution 裡 1 分人數是 0（固定 0~5 六桶，沒人選也要出現）", stat_b1.get("distribution", {}).get("1") == 0)
check("distribution 剛好 6 個桶（0~5）", len(stat_b1.get("distribution", {})) == 6)
check(
    "未作答（key 不存在）與非法值（9、not-a-number）都沒有被計入 answered_count",
    stat_b1.get("answered_count") == 3,  # 6 筆裡只有 3 筆合法
)

print("\n--- B2：完全沒人回答的 rating 題 ---")
items_b2 = [{"id": "q_empty", "type": "rating", "title": "沒人回答這題"}]
responses_b2 = [FakeResponse({}), FakeResponse({"other_q": "x"})]
stats_b2 = _build_rating_stats(items_b2, responses_b2)
check("題目仍然出現在結果陣列裡（不會因為沒人答就消失）", len(stats_b2) == 1)
check("average 是 None", stats_b2[0].get("average") is None)
check("answered_count 是 0", stats_b2[0].get("answered_count") == 0)
check("distribution 六桶全部是 0", all(v == 0 for v in stats_b2[0].get("distribution", {}).values()))

print("\n--- B3：question_title 欄位相容 ---")
items_b3 = [{"id": "q_qt", "type": "rating", "question_title": "用 question_title 欄位命名的題目"}]
stats_b3 = _build_rating_stats(items_b3, [FakeResponse({"q_qt": "4"})])
check("title 缺席時能讀到 question_title", stats_b3[0].get("title") == "用 question_title 欄位命名的題目")

print("\n--- B4：問卷完全沒有 rating 題 ---")
items_b4 = [{"id": "q1", "type": "short", "title": "只有開放題"}]
stats_b4 = _build_rating_stats(items_b4, [FakeResponse({"q1": "文字回答"})])
check("rating_stats 是空陣列", stats_b4 == [])

print("\n--- B5：沒有任何 responses ---")
stats_b5 = _build_rating_stats(items_b1, [])
check("沒有任何回覆時，rating 題仍出現、average=None、answered_count=0", (
    len(stats_b5) == 1
    and stats_b5[0]["average"] is None
    and stats_b5[0]["answered_count"] == 0
))


# ═══════════════════════════════════════════════════════════════
# Part C：POST /api/surveys/<access_code>/analyze 端到端測試
# ═══════════════════════════════════════════════════════════════
print("\n========== Part C：/api/surveys/<code>/analyze 端到端 ==========")

print("\n--- C1：rating-only 問卷（沒有 short 題，走早退分支）---")
tid_rating_only = make_survey(
    access_code="RATEONLY",
    items=[
        {"id": "q_r1", "type": "rating", "title": "滿意度"},
        {"id": "q_r2", "type": "rating", "title": "推薦意願"},
    ],
)
add_response(tid_rating_only, {"q_r1": "0", "q_r2": "5"})
add_response(tid_rating_only, {"q_r1": "3"})  # q_r2 未作答

resp_c1 = client.post("/api/surveys/RATEONLY/analyze", headers=auth_header(1))
data_c1 = resp_c1.get_json()
check("HTTP 200", resp_c1.status_code == 200)
check("aggregated_groups 是空陣列（沒有 short 題可分類）", data_c1.get("aggregated_groups") == [])
check("rating_stats 長度為 2", len(data_c1.get("rating_stats") or []) == 2)
q_r1_stat = next((s for s in data_c1["rating_stats"] if s["question_id"] == "q_r1"), None)
q_r2_stat = next((s for s in data_c1["rating_stats"] if s["question_id"] == "q_r2"), None)
check("q_r1（含 0 分）answered_count=2, average=(0+3)/2=1.5", q_r1_stat and q_r1_stat["answered_count"] == 2 and q_r1_stat["average"] == 1.5)
check("q_r1 distribution 0 分人數是 1（0 分沒被誤判成未作答）", q_r1_stat and q_r1_stat["distribution"]["0"] == 1)
check("q_r2 只有 1 筆作答（5 分），另一筆未作答不計入", q_r2_stat and q_r2_stat["answered_count"] == 1 and q_r2_stat["average"] == 5.0)
check("q_r1 是 items 第 1 題，question_number=1", q_r1_stat and q_r1_stat["question_number"] == 1)
check("q_r2 是 items 第 2 題，question_number=2", q_r2_stat and q_r2_stat["question_number"] == 2)


print("\n--- C2：short-only 問卷，Published Taxonomy 可用（既有分類 regression）---")
tid_short_only = make_survey(
    access_code="SHORTONLY",
    items=[{"id": "q1", "type": "short", "title": "對主管的建議", "question_type": QUESTION_LEADERSHIP}],
)
text_c2 = "希望主管可以多給一些工作上的回饋"
add_response(tid_short_only, {"q1": text_c2})
masked_c2 = mask_pii(text_c2)
q({"segments": [masked_c2]})
q({"classifications": [{"index": 0, "main_category": "部門合作", "sub_category": "A2 回饋與溝通",
                          "secondary_sub_category": None, "reasoning": "r", "summary": "s", "confidence": "high"}]})

resp_c2 = client.post("/api/surveys/SHORTONLY/analyze", headers=auth_header(1))
data_c2 = resp_c2.get_json()
check("HTTP 200", resp_c2.status_code == 200)
check("既有分類行為不受影響：aggregated_groups 有 1 組", len(data_c2.get("aggregated_groups") or []) == 1)
check("rating_stats 是空陣列（這份問卷沒有 rating 題）", data_c2.get("rating_stats") == [])


print("\n--- C3：rating + short 混合，taxonomy 可用，兩邊互不干擾 ---")
tid_mixed = make_survey(
    access_code="MIXED01",
    items=[
        {"id": "m_rating", "type": "rating", "title": "課程滿意度"},
        {"id": "m_short", "type": "short", "title": "改善建議", "question_type": QUESTION_LEADERSHIP},
    ],
)
text_c3 = "希望增加人力資源"
add_response(tid_mixed, {"m_rating": "0", "m_short": text_c3})
masked_c3 = mask_pii(text_c3)
q({"segments": [masked_c3]})
q({"classifications": [{"index": 0, "main_category": "部門合作", "sub_category": "B2 支援協作",
                          "secondary_sub_category": None, "reasoning": "r", "summary": "s", "confidence": "high"}]})

resp_c3 = client.post("/api/surveys/MIXED01/analyze", headers=auth_header(1))
data_c3 = resp_c3.get_json()
check("HTTP 200", resp_c3.status_code == 200)
check("short 題正常送進 Gemini 分類，aggregated_groups 有 1 組", len(data_c3.get("aggregated_groups") or []) == 1)
check("rating 題的統計正確算出（rating=0 沒被吃掉）", (
    len(data_c3.get("rating_stats") or []) == 1
    and data_c3["rating_stats"][0]["answered_count"] == 1
    and data_c3["rating_stats"][0]["average"] == 0.0
    and data_c3["rating_stats"][0]["distribution"]["0"] == 1
))
with app.app_context():
    rc_for_mixed = m.Response_Classification.query.all()
    rating_ids_in_classification = [r.question_id for r in rc_for_mixed if r.question_id == "m_rating"]
    check(
        "rating 題的 question_id 完全沒有出現在 Response_Classification 裡（沒進 Gemini/Taxonomy）",
        len(rating_ids_in_classification) == 0,
    )


print("\n--- C4：rating + short 混合，short 題 taxonomy 不可用 ---")
tid_mixed_notax = make_survey(
    access_code="MIXED02",
    items=[
        {"id": "n_rating", "type": "rating", "title": "整體評價"},
        {"id": "n_short", "type": "short", "title": "沒有對應 taxonomy 的開放題", "question_type": "career_and_feedback"},
    ],
)
add_response(tid_mixed_notax, {"n_rating": "5", "n_short": "隨便寫點什麼"})

resp_c4 = client.post("/api/surveys/MIXED02/analyze", headers=auth_header(1))
data_c4 = resp_c4.get_json()
check("HTTP 200", resp_c4.status_code == 200)
check("short 題因為 taxonomy 不可用，沒有任何分類結果", data_c4.get("aggregated_groups") == [])
check(
    "即使 short 題分類不可用，rating 統計仍然正確算出（rating=5）",
    len(data_c4.get("rating_stats") or []) == 1
    and data_c4["rating_stats"][0]["answered_count"] == 1
    and data_c4["rating_stats"][0]["average"] == 5.0,
)


# ═══════════════════════════════════════════════════════════════
# Part D：build_xlsx() / build_docx() 的 rating_stats 可選參數
# ═══════════════════════════════════════════════════════════════
print("\n========== Part D：build_xlsx / build_docx 的 rating_stats 參數 ==========")

sample_classification_rows = [
    {
        "main_category": "部門合作",
        "sub_category": "B2 支援協作",
        "respondent_text": "希望增加人力支援",
        "aggregated_reasoning": "reasoning",
        "aggregated_summary": "summary",
    },
]
sample_rating_stats = [
    {
        "question_id": "q_rating",
        "question_number": 1,
        "title": "課程整體滿意度",
        "average": 2.7,
        "answered_count": 3,
        "distribution": {"0": 1, "1": 0, "2": 0, "3": 1, "4": 0, "5": 1},
    },
]

print("\n--- D1：有 rating_stats，rows 也有內容 -> Excel 用極簡正式報表版型呈現（不使用圖表／bar／卡片底色）---")
xlsx_with_rating = build_xlsx(sample_classification_rows, title="分類結果", rating_stats=sample_rating_stats)
sheets_with_rating, sheet_names_with_rating = read_xlsx_sheets(xlsx_with_rating)
check("評分題統計 sheet 存在", "評分題統計" in sheets_with_rating)
check("分類結果 sheet 仍然存在", "分類結果" in sheets_with_rating)

_wb_check = openpyxl.load_workbook(io.BytesIO(xlsx_with_rating))
_rating_ws_check = _wb_check["評分題統計"]
check("評分題統計 sheet 完全沒有使用 Excel 圖表物件", len(_rating_ws_check._charts) == 0)
check(
    "評分題統計 sheet 完全沒有套用任何 conditional formatting（不使用 data bar／長條圖）",
    len(list(_rating_ws_check.conditional_formatting)) == 0,
)

rating_sheet_rows = sheets_with_rating["評分題統計"]
check("Q 編號＋題目全文出現在題目列（同一格，橫跨整個區塊寬度）", rating_sheet_rows[0][0] == "Q1　課程整體滿意度")
check("平均分數行格式為「平均分數：」＋「2.7 / 5」兩個儲存格", rating_sheet_rows[1][0] == "平均分數：" and rating_sheet_rows[1][1] == "2.7 / 5")
check("有效回答行格式為「有效回答：」＋「3 份」", rating_sheet_rows[2][0] == "有效回答：" and rating_sheet_rows[2][1] == "3 份")
check("分布區只用兩列：第一列是 0~5 分表頭", rating_sheet_rows[4] == ["0 分", "1 分", "2 分", "3 分", "4 分", "5 分"])
check(
    "分布區第二列是對應人數，0 分那格是「1 人」（不是被當成未作答漏掉）",
    rating_sheet_rows[5][0] == "1 人",
)
check("分布區第二列裡 5 分那格是「1 人」", rating_sheet_rows[5][5] == "1 人")

average_value_cell = _rating_ws_check.cell(row=2, column=2)
title_cell_check = _rating_ws_check.cell(row=1, column=1)
check(
    "只有平均分數這個關鍵數字用粉色強調，題目列文字是深灰色（不是整條高飽和桃紅底白字）",
    average_value_cell.font.color.rgb in ("FFF43F5E", "00F43F5E")
    and title_cell_check.font.color is not None
    and title_cell_check.font.color.rgb not in ("FFFFFFFF", "00FFFFFF"),
)
check("分類結果 sheet 內容完全沒被影響", sheets_with_rating["分類結果"][0] == COLUMN_HEADERS)

print("\n--- D1b：多題時每題各自一個獨立區塊，題目之間留 2 列空白 ---")
two_question_stats = sample_rating_stats + [
    {
        "question_id": "q_rating_2",
        "question_number": 2,
        "title": "講師表達能力",
        "average": 4.0,
        "answered_count": 2,
        "distribution": {"0": 0, "1": 0, "2": 0, "3": 0, "4": 2, "5": 0},
    },
]
xlsx_two_questions = build_xlsx([], title="分類結果", rating_stats=two_question_stats)
sheets_two_questions, _ = read_xlsx_sheets(xlsx_two_questions)
two_q_rows = sheets_two_questions["評分題統計"]
check(
    "第一題區塊佔用第 1~6 列（題目、平均分、有效回答、空白、分布表頭、分布數值）",
    two_q_rows[0][0] == "Q1　課程整體滿意度" and two_q_rows[5][0] == "1 人",
)
check("第 7、8 列是空白列（題目之間留白，不套用任何內容）", two_q_rows[6] == [None] * 6 and two_q_rows[7] == [None] * 6)
check("第二題區塊從第 9 列重新開始，兩題不會黏在一起", two_q_rows[8][0] == "Q2　講師表達能力")

print("\n--- D2：有 rating_stats，但 rows=[]（問卷只有 rating 題）不能壞 ---")
try:
    xlsx_rating_only = build_xlsx([], title="分類結果", rating_stats=sample_rating_stats)
    sheets_rating_only, _ = read_xlsx_sheets(xlsx_rating_only)
    check("rows=[] 時 build_xlsx 不拋例外，評分題統計 sheet 仍正常產生", "評分題統計" in sheets_rating_only)
    check(
        "rows=[] 時評分題統計內容仍然正確（不因為分類結果是空的而跟著壞掉）",
        sheets_rating_only["評分題統計"][1][1] == "2.7 / 5",
    )
except Exception as e:
    check(f"rows=[] 時 build_xlsx 不拋例外（實際拋出：{e!r}）", False)

print("\n--- D3：沒有 rating_stats（None）-> Excel 輸出跟新增這個參數之前完全一樣（regression）---")
xlsx_no_rating = build_xlsx(sample_classification_rows, title="分類結果")
sheets_no_rating, sheet_names_no_rating = read_xlsx_sheets(xlsx_no_rating)
check("沒有評分題統計 sheet", "評分題統計" not in sheets_no_rating)
check("只有 1 張 sheet（跟這個參數新增之前一樣）", len(sheet_names_no_rating) == 1)
check(
    "分類結果 sheet 內容跟有 rating_stats 版本一致（同一份 rows 不因為多一個參數而變）",
    sheets_no_rating["分類結果"] == sheets_with_rating["分類結果"],
)

print("\n--- D4：Word 有 rating_stats -> 每題一個區塊，插入甜甜圈圖圖片，不是表格或純文字清單 ---")
docx_with_rating = build_docx(sample_classification_rows, title="分類結果", rating_stats=sample_rating_stats)
docx_text_with_rating = read_docx_text(docx_with_rating)
check("Word 裡出現「評分題統計」標題", "評分題統計" in docx_text_with_rating)
check("Word 裡題目全文清楚出現，作為小標題（Q 編號 + 題目同一行）", "Q1　課程整體滿意度" in docx_text_with_rating)
check("Word 裡平均分視覺突出的文字「2.7 / 5」有輸出", "2.7 / 5" in docx_text_with_rating)
check("Word 裡有效回答數正確輸出", "有效回答：3 份" in docx_text_with_rating)
check("Word 裡精簡列出 0 分：1 人（0 分沒被當成未作答漏掉）", "0 分：1 人" in docx_text_with_rating)
check("Word 裡精簡列出 5 分：1 人", "5 分：1 人" in docx_text_with_rating)
check(
    "顯示順序：評分題統計在既有分類結果標題之前",
    docx_text_with_rating.index("評分題統計") < docx_text_with_rating.index("部門合作"),
)

_doc_check = _DocxDocument(io.BytesIO(docx_with_rating))
check(
    "評分題統計不是用表格呈現：整份文件裡只有 1 張表格（既有分類結果那張），\n"
    "    評分題統計那段完全沒有產生表格",
    len(_doc_check.tables) == 1,
)
_inline_shapes = _doc_check.inline_shapes
check(
    "評分題統計改成插入甜甜圈圖圖片：文件裡有 1 張內嵌圖片（對應這 1 題）",
    len(_inline_shapes) == 1,
)

print("\n--- D5：Word 有 rating_stats，但 rows=[]（rating-only）不能壞 ---")
try:
    docx_rating_only = build_docx([], title="問卷分析", rating_stats=sample_rating_stats)
    docx_text_rating_only = read_docx_text(docx_rating_only)
    check("rows=[] 時 build_docx 不拋例外", "評分題統計" in docx_text_rating_only)
    check("rows=[] 時評分題統計內容仍正確（平均分 2.7 / 5 仍有輸出）", "2.7 / 5" in docx_text_rating_only)
    _doc_rating_only_check = _DocxDocument(io.BytesIO(docx_rating_only))
    check("rows=[] 時甜甜圈圖圖片仍正常插入", len(_doc_rating_only_check.inline_shapes) == 1)
except Exception as e:
    check(f"rows=[] 時 build_docx 不拋例外（實際拋出：{e!r}）", False)

print("\n--- D6：沒有 rating_stats（None）-> Word 輸出跟新增這個參數之前完全一樣（regression）---")
docx_no_rating = build_docx(sample_classification_rows, title="分類結果")
docx_text_no_rating = read_docx_text(docx_no_rating)
check("Word 裡沒有「評分題統計」字樣", "評分題統計" not in docx_text_no_rating)
check("Word 裡既有分類內容仍然存在", "部門合作" in docx_text_no_rating and "B2 支援協作" in docx_text_no_rating)
check("沒有 rating_stats 時完全不插入任何圖片", len(_DocxDocument(io.BytesIO(docx_no_rating)).inline_shapes) == 0)


# ═══════════════════════════════════════════════════════════════
# Part F：_render_rating_donut_png() 圖片產生器單元測試
# ═══════════════════════════════════════════════════════════════
print("\n========== Part F：_render_rating_donut_png() 單元測試 ==========")

png_normal = _render_rating_donut_png({"0": 1, "1": 0, "2": 0, "3": 1, "4": 0, "5": 1}, 2.7)
check("正常情況（有資料）能產生非空的 PNG bytes", isinstance(png_normal, (bytes, bytearray)) and len(png_normal) > 0)
check("PNG 檔頭正確（合法圖片檔案）", png_normal[:8] == b"\x89PNG\r\n\x1a\n")

png_no_data = _render_rating_donut_png({"0": 0, "1": 0, "2": 0, "3": 0, "4": 0, "5": 0}, None)
check(
    "完全沒有人作答（average=None、distribution 全 0）時，圖片仍能正常產生，不拋例外",
    isinstance(png_no_data, (bytes, bytearray)) and len(png_no_data) > 0,
)

png_rating_zero_only = _render_rating_donut_png({"0": 5, "1": 0, "2": 0, "3": 0, "4": 0, "5": 0}, 0.0)
check(
    "全部受試者都回答 0 分（average=0.0）時仍能正常產生圖片，不會因為 0 是 falsy 而被當成沒資料",
    isinstance(png_rating_zero_only, (bytes, bytearray)) and len(png_rating_zero_only) > 0,
)


# ═══════════════════════════════════════════════════════════════
# Part E：POST /api/exports，rating-only（rows=[] + rating_stats 非空）
# ═══════════════════════════════════════════════════════════════
print("\n========== Part E：/api/exports rating-only 情境 ==========")

with app.app_context():
    workspace = m.Workspace(user_id=1, project_name="評分統計測試工作區")
    db.session.add(workspace)
    db.session.flush()
    chat = m.Chat_History(project_id=workspace.project_id, message_content="x", sender_type="ai")
    db.session.add(chat)
    db.session.commit()
    chat_id = chat.chat_id

resp_e1 = client.post(
    "/api/exports",
    headers={"Content-Type": "application/json", **auth_header(1)},
    json={
        "chat_id": chat_id,
        "filename": "評分統計_問卷分析.xlsx",
        "export_type": "xlsx",
        "row_count": 0,
        "rows": [],
        "rating_stats": sample_rating_stats,
        "title": "問卷分析",
    },
)
check("rows=[] 但 rating_stats 非空 -> 201 成功（不因為 rows 是空陣列被擋掉）", resp_e1.status_code == 201)
export_id_e1 = (resp_e1.get_json() or {}).get("export_id")

if export_id_e1:
    resp_e1_download = client.get(f"/api/exports/{export_id_e1}/download", headers=auth_header(1))
    check("下載成功", resp_e1_download.status_code == 200)
    sheets_e1, _ = read_xlsx_sheets(resp_e1_download.data)
    check("下載的檔案裡有評分題統計 sheet", "評分題統計" in sheets_e1)

print("\n--- E2：rows=[] 且沒有 rating_stats -> 仍維持既有行為，400 ---")
resp_e2 = client.post(
    "/api/exports",
    headers={"Content-Type": "application/json", **auth_header(1)},
    json={
        "chat_id": chat_id,
        "filename": "空匯出.xlsx",
        "export_type": "xlsx",
        "row_count": 0,
        "rows": [],
        "title": "空匯出",
    },
)
check("rows=[] 且沒有 rating_stats -> 400（regression：既有防呆沒被破壞）", resp_e2.status_code == 400)


# ═══════════════════════════════════════════════════════════════
# 總結
# ═══════════════════════════════════════════════════════════════
print("\n========== 測試總結 ==========")
if FAILED:
    print(f"共有 {len(FAILED)} 項測試失敗：")
    for label in FAILED:
        print(f"  - {label}")
    sys.exit(1)
else:
    print("全部測試通過！")
